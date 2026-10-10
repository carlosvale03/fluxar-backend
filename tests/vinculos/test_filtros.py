"""
Filtros de vínculo na lista e nas exportações (VINCULO-20, VINCULO-29,
VINCULO-30, VINCULO-32 e VINCULO-33) e o filtro de valor da busca da
principal (VINCULO-03).
"""
import io
from datetime import date, timedelta

from django.core.cache import cache
from openpyxl import load_workbook

from core import travas
from tests.permissoes.base import fechar, liberacao_de_testes
from transactions.vinculos import vincular

from .base import VinculosTestCase, compra, despesa

URL_LISTA = '/api/transactions/'
URL_XLSX = '/api/export/transactions/xls/'
URL_PDF = '/api/export/transactions/pdf/'
BLOQUEADO = {
    'detail': 'Este recurso não está disponível no seu plano.', 'code': 'plan_locked', 'feature': 'vinculos',
}
ID_INEXISTENTE = '00000000-0000-4000-8000-000000000000'


def planilha(resposta):
    folha = load_workbook(io.BytesIO(b''.join(resposta.streaming_content) if resposta.streaming
                                     else resposta.content)).active
    linhas = [list(linha) for linha in folha.iter_rows(values_only=True)]
    return linhas[0], linhas[1:]


class FiltrosTestCase(VinculosTestCase):

    def setUp(self):
        super().setUp()
        cache.clear()
        travas.invalidar()
        self.addCleanup(travas.invalidar)
        self.cinema = despesa(self.a, 'Cinema', '50.00', categoria=self.a.cinema)
        self.transporte = despesa(self.a, 'Transporte', '45.00', categoria=self.a.transporte)
        self.mercado = despesa(self.a, 'Mercado', '200.00', categoria=self.a.alimentacao)
        vincular(self.transporte, self.cinema)

    def ids(self, parametros):
        resposta = self.client.get(URL_LISTA, {'page_size': 100, **parametros})
        self.assertEqual(resposta.status_code, 200, resposta.data)
        return {linha['id'] for linha in resposta.data['results']}


class PrincipalIdTests(FiltrosTestCase):

    def test_principal_id_do_cinema_traz_o_cinema_e_o_transporte(self):
        self.assertEqual(self.ids({'principalId': str(self.cinema.pk)}), {str(self.cinema.pk), str(self.transporte.pk)})

    def test_principal_id_traz_as_parcelas_das_compras(self):
        show = compra(self.a, 'Show', '300.00', parcelas=3, categoria=self.a.lazer)
        jantar = compra(self.a, 'Jantar', '120.00', parcelas=3, categoria=self.a.alimentacao)
        vincular(jantar[2], show[1])
        esperadas = {str(t.pk) for t in [*show, *jantar]}
        # Qualquer parcela da principal serve de filtro
        self.assertEqual(self.ids({'principalId': str(show[2].pk)}), esperadas)
        self.assertEqual(self.ids({'principalId': str(show[0].pk)}), esperadas)

    def test_principal_id_de_outro_usuario_inexistente_ou_malformado_recebe_400(self):
        alheia = despesa(self.b, 'Cinema B', '50.00')
        for valor in (str(alheia.pk), ID_INEXISTENTE, 'abc'):
            with self.subTest(valor=valor):
                resposta = self.client.get(URL_LISTA, {'principalId': valor})
                self.assertEqual(resposta.status_code, 400)
                self.assertEqual(resposta.data, {'principalId': [f'Transação inválida no filtro: {valor}.']})


class LinkedTests(FiltrosTestCase):

    def test_linked_traz_so_principais_e_dependentes_com_as_parcelas(self):
        show = compra(self.a, 'Show', '300.00', parcelas=3, categoria=self.a.lazer)
        jantar = compra(self.a, 'Jantar', '120.00', parcelas=3, categoria=self.a.alimentacao)
        avulsa = compra(self.a, 'Livro', '90.00', parcelas=3)
        vincular(jantar[1], show[2])

        esperadas = {str(t.pk) for t in [self.cinema, self.transporte, *show, *jantar]}
        obtidas = self.ids({'linked': 'true'})
        self.assertEqual(obtidas, esperadas)
        self.assertNotIn(str(self.mercado.pk), obtidas)
        self.assertFalse({str(t.pk) for t in avulsa} & obtidas)

    def test_linked_com_outro_valor_recebe_400_no_campo(self):
        resposta = self.client.get(URL_LISTA, {'linked': 'false'})
        self.assertEqual(resposta.status_code, 400)
        self.assertEqual(list(resposta.data), ['linked'])


class ValorTests(FiltrosTestCase):
    """VINCULO-03: a busca da principal pelo valor."""

    def test_valor_acha_uma_despesa_antiga_atras_de_mais_de_100_recentes(self):
        antiga = despesa(self.a, 'Teatro', '37.90', dia=date(2025, 1, 10))
        inicio = date(2026, 9, 1)
        for i in range(120):
            despesa(self.a, f'Café {i}', '8.00', dia=inicio + timedelta(days=i % 30))
        # Sem o filtro, a despesa antiga fica fora das 100 mais recentes
        self.assertNotIn(str(antiga.pk), self.ids({'type': 'EXPENSE'}))

        self.assertEqual(self.ids({'type': 'EXPENSE', 'amount': '37.90'}), {str(antiga.pk)})

    def test_valor_invalido_recebe_400_no_campo(self):
        for valor in ('abc', '0', '-5.00', '1.234'):
            with self.subTest(valor=valor):
                resposta = self.client.get(URL_LISTA, {'amount': valor})
                self.assertEqual(resposta.status_code, 400)
                self.assertEqual(resposta.data, {'amount': [f'Valor inválido no filtro: {valor}.']})
        self.assertEqual(self.client.get(URL_XLSX, {'amount': 'abc'}).status_code, 400)

    def test_compra_parcelada_aparece_pelo_total(self):
        show = compra(self.a, 'Show', '300.00', parcelas=3, categoria=self.a.lazer)
        compra(self.b, 'Show B', '300.00', parcelas=3)
        parcelas = {str(t.pk) for t in show}

        self.assertEqual(self.ids({'type': 'EXPENSE', 'amount': '300.00'}), parcelas)
        # Pelo valor de uma parcela também
        self.assertEqual(self.ids({'type': 'EXPENSE', 'amount': '100.00'}), parcelas)
        self.assertEqual(self.ids({'type': 'EXPENSE', 'amount': '50.00'}), {str(self.cinema.pk)})


class ExportacaoTests(FiltrosTestCase):
    """VINCULO-33."""

    def test_exportacao_com_o_filtro_do_cinema_traz_as_duas_linhas_sem_coluna_de_vinculo(self):
        sem_filtro = self.client.get(URL_XLSX)
        self.assertEqual(sem_filtro.status_code, 200)
        cabecalho_sem_filtro, _ = planilha(sem_filtro)

        resposta = self.client.get(URL_XLSX, {'principalId': str(self.cinema.pk)})

        self.assertEqual(resposta.status_code, 200)
        cabecalho, linhas = planilha(resposta)
        self.assertEqual(cabecalho, cabecalho_sem_filtro)
        coluna = cabecalho.index('Descrição')
        self.assertEqual(sorted(linha[coluna] for linha in linhas), ['Cinema', 'Transporte'])
        for titulo in cabecalho:
            self.assertNotIn('principal', str(titulo).lower())
            self.assertNotIn('víncul', str(titulo).lower())

        ligadas = self.client.get(URL_XLSX, {'linked': 'true'})
        self.assertEqual(ligadas.status_code, 200)
        self.assertEqual(len(planilha(ligadas)[1]), 2)
        self.assertEqual(self.client.get(URL_PDF, {'principalId': str(self.cinema.pk)}).status_code, 200)

    def test_exportacao_com_id_invalido_recebe_400(self):
        resposta = self.client.get(URL_XLSX, {'principalId': ID_INEXISTENTE})
        self.assertEqual(resposta.status_code, 400)
        self.assertEqual(resposta.data, {'principalId': [f'Transação inválida no filtro: {ID_INEXISTENTE}.']})


class TravaTests(FiltrosTestCase):
    """VINCULO-20."""

    def test_com_o_recurso_travado_os_filtros_recebem_403(self):
        liberacao_de_testes(False)
        fechar('vinculos')
        chamadas = [
            (URL_LISTA, {'principalId': str(self.cinema.pk)}),
            (URL_LISTA, {'linked': 'true'}),
            (URL_XLSX, {'principalId': str(self.cinema.pk)}),
            (URL_PDF, {'linked': 'true'}),
        ]
        for url, parametros in chamadas:
            with self.subTest(url=url, parametros=parametros):
                resposta = self.client.get(url, parametros)
                self.assertEqual(resposta.status_code, 403)
                self.assertEqual(resposta.data, BLOQUEADO)
        # Sem os filtros, a lista segue liberada e mostra os vínculos existentes
        resposta = self.client.get(URL_LISTA)
        self.assertEqual(resposta.status_code, 200)
        [cinema] = [linha for linha in resposta.data['results'] if linha['id'] == str(self.cinema.pk)]
        self.assertEqual(cinema['total_cost'], '95.00')
