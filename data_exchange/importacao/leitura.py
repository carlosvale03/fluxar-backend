"""
Leitura do arquivo de importação (IMPORT-01 a IMPORT-08, IMPORT-11,
IMPORT-14, IMPORT-32).

Confere a extensão e o tamanho, lê OFX, CSV ou XLSX e devolve as linhas
brutas com o número de cada uma no arquivo. Nada é gravado aqui. Uma falha
do arquivo inteiro levanta `ValidationError` no campo `file` (HTTP 400).
"""
import csv
import io
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path

import pandas as pd
from ofxparse import OfxParser
from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError

FORMATO_NAO_SUPORTADO = 'Formato não suportado. Envie um arquivo OFX, CSV ou XLSX.'
ARQUIVO_GRANDE_DEMAIS = 'O arquivo passa do limite de 5 MB.'
LINHAS_DEMAIS = 'O arquivo passa do limite de 10.000 linhas.'
ARQUIVO_ILEGIVEL = 'Não foi possível ler o arquivo.'
COLUNAS_NAO_ENCONTRADAS = 'Colunas não encontradas no arquivo: {colunas}.'
MAIS_DE_UMA_CONTA = 'O arquivo tem mais de uma conta. Exporte um arquivo por conta.'
COLUNAS_OBRIGATORIAS = 'Colunas obrigatórias não informadas para o tipo de importação.'

LIMITE_DE_BYTES = 5 * 1024 * 1024
LIMITE_DE_LINHAS = 10_000

# Colunas do mapeamento lidas em cada tipo de importação; as obrigatórias
# precisam estar mapeadas. Toda coluna mapeada precisa existir no arquivo.
COLUNAS_DO_TIPO = {
    'INCOME_EXPENSE': (
        'date_column', 'amount_column', 'description_column', 'type_column', 'status_column',
        'category_column', 'subcategory_column', 'tags_column', 'account_column',
    ),
    'TRANSFER': (
        'date_column', 'amount_column', 'source_account_column', 'dest_account_column', 'tags_column',
    ),
}
OBRIGATORIAS_DO_TIPO = {
    'INCOME_EXPENSE': ('date_column', 'amount_column', 'description_column'),
    'TRANSFER': ('date_column', 'amount_column', 'source_account_column', 'dest_account_column'),
}


@dataclass(frozen=True)
class LinhaBruta:
    """Uma linha da planilha: o número dela no arquivo (cabeçalho = 1) e os valores por coluna."""
    numero: int
    valores: dict


@dataclass(frozen=True)
class TransacaoOfx:
    """Uma transação do OFX, com a posição dela no extrato a partir de 1."""
    numero: int
    data: date
    valor: Decimal
    memo: str
    favorecido: str
    fitid: str | None


def recusar(mensagem):
    raise ValidationError({'file': [mensagem]})


def tipo_da_importacao(import_type):
    return 'TRANSFER' if import_type == 'TRANSFER' else 'INCOME_EXPENSE'


def conferir_arquivo(arquivo, extensoes):
    """Extensão entre `extensoes` (IMPORT-02) e até 5 MB, antes de ler (IMPORT-04)."""
    if Path(arquivo.name or '').suffix.lower() not in extensoes:
        recusar(FORMATO_NAO_SUPORTADO)
    if arquivo.size > LIMITE_DE_BYTES:
        recusar(ARQUIVO_GRANDE_DEMAIS)
    return Path(arquivo.name).suffix.lower()


def vazio(valor):
    return valor is None or (isinstance(valor, str) and not valor.strip())


def colunas_mapeadas(mapeamento, import_type):
    """As colunas do arquivo que o mapeamento usa, na ordem dos campos e sem repetir."""
    tipo = tipo_da_importacao(import_type)
    if not all(mapeamento.get(campo) for campo in OBRIGATORIAS_DO_TIPO[tipo]):
        raise ValidationError({'mapping': [COLUNAS_OBRIGATORIAS]})
    colunas = []
    for campo in COLUNAS_DO_TIPO[tipo]:
        coluna = mapeamento.get(campo)
        if coluna and coluna not in colunas:
            colunas.append(coluna)
    return colunas


def conferir_colunas(cabecalho, colunas):
    """Recusa o arquivo sem alguma coluna mapeada, citando as que faltam (IMPORT-06)."""
    faltando = [c for c in colunas if c not in cabecalho]
    if faltando:
        recusar(COLUNAS_NAO_ENCONTRADAS.format(colunas=', '.join(faltando)))


def ler_planilha(arquivo, mapeamento, import_type):
    """
    Lê um CSV ou XLSX e devolve as linhas de dados não vazias como
    `LinhaBruta` (IMPORT-30, IMPORT-32). Mais de 10.000 linhas de dados
    recusa o arquivo (IMPORT-05).
    """
    extensao = conferir_arquivo(arquivo, {'.csv', '.xlsx'})
    linhas = ler_csv(arquivo) if extensao == '.csv' else ler_xlsx(arquivo)

    cabecalho = next(linhas)
    conferir_colunas(cabecalho, colunas_mapeadas(mapeamento, import_type))
    resultado = []
    for numero, valores in linhas:
        if all(vazio(v) for v in valores):
            continue
        if len(resultado) == LIMITE_DE_LINHAS:
            recusar(LINHAS_DEMAIS)
        resultado.append(LinhaBruta(numero, dict(zip(cabecalho, valores))))
    return resultado


def ler_csv(arquivo):
    """
    O cabeçalho e depois `(número, valores)` de cada linha, tudo como texto
    para a interpretação aplicar AD-009. Codificação UTF-8 (com ou sem BOM)
    ou Windows-1252, separador `;` ou `,` pela primeira linha (IMPORT-08).
    As linhas vazias continuam na leitura para não deslocar a numeração.
    """
    conteudo = arquivo.read()
    try:
        texto = conteudo.decode('utf-8-sig')
    except UnicodeDecodeError:
        try:
            texto = conteudo.decode('cp1252')
        except UnicodeDecodeError:
            recusar(ARQUIVO_ILEGIVEL)
    primeira_linha = texto.split('\n', 1)[0]
    try:
        separador = csv.Sniffer().sniff(primeira_linha, delimiters=';,').delimiter
    except csv.Error:
        separador = ','
    try:
        df = pd.read_csv(
            io.StringIO(texto), sep=separador, dtype=str,
            keep_default_na=False, skip_blank_lines=False,
        )
    except Exception:  # noqa: BLE001 - qualquer falha do pandas é arquivo ilegível
        recusar(ARQUIVO_ILEGIVEL)

    yield list(df.columns)
    for indice, valores in enumerate(df.itertuples(index=False, name=None)):
        # Campo ausente vem como NaN do pandas
        yield indice + 2, [None if isinstance(v, float) else v for v in valores]


def ler_xlsx(arquivo):
    """
    O cabeçalho e depois `(número, valores)` de cada linha da primeira aba.
    Células numéricas e de data mantêm o tipo (IMPORT-11, IMPORT-14).
    """
    try:
        planilha = load_workbook(io.BytesIO(arquivo.read()), read_only=True, data_only=True)
        linhas = list(planilha.worksheets[0].iter_rows(values_only=True))
    except Exception:  # noqa: BLE001 - arquivo que o openpyxl não abre é ilegível
        recusar(ARQUIVO_ILEGIVEL)
    if not linhas:
        recusar(ARQUIVO_ILEGIVEL)

    yield [str(c) if c is not None else '' for c in linhas[0]]
    for numero, valores in enumerate(linhas[1:], start=2):
        yield numero, list(valores)


def ler_ofx(arquivo):
    """
    Lê um OFX de uma única conta (IMPORT-07) e devolve as transações com a
    posição de cada uma no extrato. Mais de 10.000 transações recusa o
    arquivo (IMPORT-05).
    """
    conferir_arquivo(arquivo, {'.ofx'})
    try:
        ofx = OfxParser.parse(io.BytesIO(arquivo.read()))
    except Exception:  # noqa: BLE001 - OFX que o ofxparse não lê é ilegível
        recusar(ARQUIVO_ILEGIVEL)
    if len(ofx.accounts) > 1:
        recusar(MAIS_DE_UMA_CONTA)
    if not ofx.accounts or ofx.accounts[0].statement is None:
        recusar(ARQUIVO_ILEGIVEL)

    transacoes = ofx.accounts[0].statement.transactions
    if len(transacoes) > LIMITE_DE_LINHAS:
        recusar(LINHAS_DEMAIS)
    return [
        TransacaoOfx(
            numero=numero, data=tx.date.date(), valor=Decimal(str(tx.amount)),
            memo=(tx.memo or '').strip(), favorecido=(tx.payee or '').strip(),
            fitid=(tx.id or '').strip() or None,
        )
        for numero, tx in enumerate(transacoes, start=1)
    ]


def contas_da_planilha(linhas, mapeamento, import_type):
    """Os nomes de conta da planilha, sem repetir e em ordem, para a etapa de mapeamento."""
    if tipo_da_importacao(import_type) == 'TRANSFER':
        campos = ('source_account_column', 'dest_account_column')
    else:
        campos = ('account_column',)
    colunas = [mapeamento[c] for c in campos if mapeamento.get(c)]
    nomes = {
        str(linha.valores.get(coluna)).strip()
        for linha in linhas for coluna in colunas
        if not vazio(linha.valores.get(coluna))
    }
    return sorted(nomes)
