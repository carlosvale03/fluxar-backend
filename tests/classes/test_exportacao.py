"""
Coluna "Classe" na exportação em XLSX (CLASSE-40): a das transações e o
download dos dados da LGPD. A usuária tem o padrão do cadastro (CLASSE-24).
"""
import io
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from accounts.models import Account, CreditCard
from transactions.models import Category, Transaction
from transactions.services import TransactionService

from .base import ClassesTestCase, categoria

URL_XLS = '/api/export/transactions/xls/'
URL_MEUS_DADOS = '/api/users/me/export/'


def linhas_por_descricao(conteudo, aba=None):
    """`{descrição: {coluna: valor}}` da aba, com as células vazias como ""."""
    planilha = load_workbook(io.BytesIO(conteudo))
    folha = planilha[aba] if aba else planilha.active
    linhas = list(folha.iter_rows(values_only=True))
    cabecalho = list(linhas[0])
    resultado = {}
    for linha in linhas[1:]:
        valores = {coluna: ('' if valor is None else valor) for coluna, valor in zip(cabecalho, linha)}
        resultado[valores['Descrição']] = valores
    return cabecalho, resultado


class ColunaClasseTests(ClassesTestCase):

    def setUp(self):
        super().setUp()
        ana = self.ana
        self.conta = Account.objects.create(user=ana, name='Corrente', type='CHECKING', initial_balance=Decimal('5000'))
        poupanca = Account.objects.create(user=ana, name='Poupança', type='SAVINGS')
        cartao = CreditCard.objects.create(
            user=ana, name='Cartão', limit=Decimal('10000.00'), closing_day=5, due_day=15, account=self.conta,
        )
        comida = Category.objects.get(user=ana, name='Comida')
        dia = date(2026, 9, 10)
        self.lancar('Doce', Category.objects.get(user=ana, name='Doces'), dia)
        self.lancar('Pedido', categoria(ana, 'Delivery', parent=comida), dia)
        self.lancar('No mercado', categoria(ana, 'Mercado'), dia)
        self.lancar('Sem categoria', None, dia)
        self.lancar('Salário', Category.objects.get(user=ana, name='Salário'), dia, tipo='INCOME')
        TransactionService.create_credit_card_expense(
            ana, cartao, Decimal('50.00'), dia, 'Almoço no cartão', comida, installments=1,
        )
        TransactionService.create_credit_card_expense(
            ana, cartao, Decimal('20.00'), dia, 'Cartão sem categoria', None, installments=1,
        )
        TransactionService.create_transfer(ana, self.conta, poupanca, Decimal('100.00'), dia)

    def lancar(self, descricao, categoria_, dia, tipo='EXPENSE'):
        return Transaction.objects.create(
            user=self.ana, type=tipo, status='COMPLETED', amount=Decimal('10.00'), date=dia,
            account=self.conta, category=categoria_, description=descricao,
        )

    def classes(self, linhas):
        return {descricao: valores['Classe'] for descricao, valores in linhas.items()}

    def test_xlsx_das_transacoes_tem_a_classe_efetiva_das_despesas_e_compras_no_cartao(self):
        resposta = self.client.get(URL_XLS)

        self.assertEqual(resposta.status_code, 200)
        cabecalho, linhas = linhas_por_descricao(resposta.content)
        self.assertEqual(cabecalho[-1], 'Classe')
        sem_transferencias = {
            d: linha for d, linha in linhas.items() if not linha['Tipo'].startswith('Transferência')
        }
        self.assertEqual(self.classes(sem_transferencias), {
            'Doce': 'Dispensável',
            'Pedido': 'Essencial',
            'No mercado': 'Sem classe',
            'Sem categoria': 'Sem classe',
            'Salário': '',
            'Almoço no cartão': 'Essencial',
            'Cartão sem categoria': 'Sem classe',
        })
        # As pernas da transferência ficam sem classe
        planilha = load_workbook(io.BytesIO(resposta.content)).active
        todas = list(planilha.iter_rows(values_only=True))
        tipo, classe = cabecalho.index('Tipo'), cabecalho.index('Classe')
        transferencias = [linha for linha in todas[1:] if str(linha[tipo]).startswith('Transferência')]
        self.assertEqual(len(transferencias), 2)
        self.assertEqual({linha[classe] or '' for linha in transferencias}, {''})

    def test_xlsx_com_o_filtro_de_classe_traz_as_mesmas_linhas_com_a_coluna(self):
        resposta = self.client.get(URL_XLS, {'classId': str(self.essencial().pk)})

        self.assertEqual(resposta.status_code, 200)
        _, linhas = linhas_por_descricao(resposta.content)
        self.assertEqual(self.classes(linhas), {'Pedido': 'Essencial', 'Almoço no cartão': 'Essencial'})

    def test_xlsx_dos_dados_da_lgpd_tambem_tem_a_coluna(self):
        resposta = self.client.get(URL_MEUS_DADOS)

        self.assertEqual(resposta.status_code, 200)
        cabecalho, linhas = linhas_por_descricao(resposta.content, 'Transações')
        self.assertEqual(cabecalho, [
            'Data', 'Descrição', 'Valor', 'Conta', 'Situação', 'Categoria', 'Subcategoria', 'Tags', 'Tipo', 'Classe',
        ])
        classes = self.classes(linhas)
        self.assertEqual(classes['Doce'], 'Dispensável')
        self.assertEqual(classes['Pedido'], 'Essencial')
        self.assertEqual(classes['Cartão sem categoria'], 'Sem classe')
        self.assertEqual(classes['Salário'], '')
