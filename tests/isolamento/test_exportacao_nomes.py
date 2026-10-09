"""
Transações de A ligadas à força à conta, à categoria, à categoria-pai e à
tag de B, como antes da correção: a exportação em planilha e em PDF não
escreve nenhum nome de B, e trata essas relações como ausentes (ISOL-15).
"""
import base64
import io
import re
import zlib
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from transactions.models import Category, Transaction
from tests.isolamento.base import DoisUsuariosTestCase

NOMES_DE_B = ('Conta B', 'Mercado B', 'Feira B', 'Viagem B')


def texto_do_pdf(conteudo):
    """
    Texto das páginas do PDF gerado pelo reportlab, que grava cada página em
    ASCII85 sobre zlib.
    """
    paginas = re.findall(rb'stream\r?\n(.*?)endstream', conteudo, re.S)
    return b''.join(zlib.decompress(base64.a85decode(p.strip(), adobe=True)) for p in paginas).decode('latin-1')


class ExportacaoSemNomesDeBTests(DoisUsuariosTestCase):

    def setUp(self):
        pendurada = Category.objects.create(
            user=self.a.usuario, name='Pendurada A', type='EXPENSE', parent=self.b.categoria,
        )
        # Todas as relações apontam para B, com uma tag própria de A junto
        self.transacao('Forjada', 5, self.b.conta, self.b.subcategoria, [self.a.tag, self.b.tag])
        # Categoria própria de A pendurada na categoria de B
        self.transacao('Pendurada', 6, self.a.conta, pendurada, [self.b.tag])
        self.transacao('Própria', 7, self.a.conta, self.a.subcategoria, [self.a.tag])
        self.cliente = self.como(self.a.usuario)

    def transacao(self, descricao, dia, conta, categoria, tags):
        t = Transaction.objects.create(
            user=self.a.usuario, type='EXPENSE', status='COMPLETED', description=descricao,
            amount=Decimal('10.00'), date=date(2026, 9, dia), account=conta, category=categoria,
        )
        t.tags.set(tags)

    def exportar(self, formato):
        resp = self.cliente.get(f'/api/export/transactions/{formato}/')
        self.assertEqual(resp.status_code, 200)
        return resp.content

    def test_planilha_nao_tem_nomes_de_b(self):
        planilha = load_workbook(io.BytesIO(self.exportar('xls'))).active
        cabecalho, *linhas = [[c or '' for c in linha] for linha in planilha.iter_rows(values_only=True)]
        colunas = [cabecalho.index(nome) for nome in ('Descrição', 'Conta', 'Categoria', 'Subcategoria', 'Tags')]

        self.assertEqual([tuple(linha[i] for i in colunas) for linha in linhas], [
            ('Forjada', '', '', '', 'Viagem A'),
            ('Pendurada', 'Conta A', 'Pendurada A', '', ''),
            ('Própria', 'Conta A', 'Mercado A', 'Feira A', 'Viagem A'),
        ])
        texto = str(linhas)
        for nome in NOMES_DE_B:
            self.assertNotIn(nome, texto)

    def test_pdf_nao_tem_nomes_de_b(self):
        texto = texto_do_pdf(self.exportar('pdf'))

        for nome in NOMES_DE_B:
            self.assertNotIn(nome, texto)
        # A linha forjada sai com a conta em branco (REL-25) e "-" na categoria
        self.assertEqual(texto.count('Tm  T* ET'), 1)
        self.assertEqual(texto.count('(-) Tj'), 1)
        for nome in ('(Conta A)', '(Pendurada A)', '(Feira A)'):
            self.assertIn(nome, texto)
