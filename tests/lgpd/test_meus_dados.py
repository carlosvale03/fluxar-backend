"""
Download dos dados financeiros antes da exclusão (LGPD-02).

O XLSX sai em qualquer plano, mesmo com a exportação de transações travada,
com uma aba por tipo de dado e só os dados do próprio usuário.
"""
import io
from datetime import date
from decimal import Decimal

from django.core.cache import cache
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from api.models import GlobalSetting, TravaDePlano
from core import travas
from tests.isolamento.base import criar_dados
from transactions.models import Transaction

EXPORTAR = '/api/users/me/export/'
EXPORTAR_TRANSACOES = '/api/export/transactions/xls/'
ABAS = ['Transações', 'Contas', 'Cartões', 'Metas', 'Orçamentos']


class MeusDadosTests(APITestCase):

    def setUp(self):
        cache.clear()
        travas.invalidar()
        self.addCleanup(travas.invalidar)
        self.a = criar_dados('A')
        self.b = criar_dados('B')
        for dono in (self.a, self.b):
            Transaction.objects.create(
                user=dono.usuario, account=dono.conta, category=dono.categoria, type='EXPENSE',
                status='COMPLETED', description=f'Feira do mês {dono.usuario.name}',
                amount=Decimal('87.50'), date=date(2026, 9, 15),
            )
        # A exportação de transações em XLSX travada no plano Comum, sem a
        # liberação para testes
        GlobalSetting.objects.update_or_create(key='testing_unlock', defaults={'value': 'false'})
        TravaDePlano.objects.create(chave='exportacao_xlsx', plano='COMMON', liberado=False)
        travas.invalidar()
        self.client.force_authenticate(user=self.a.usuario)

    def planilha(self, resposta):
        return load_workbook(io.BytesIO(resposta.content))

    def textos(self, aba):
        return {str(celula) for linha in aba.iter_rows(values_only=True) for celula in linha if celula is not None}

    def test_usuario_comum_com_a_exportacao_travada_baixa_o_xlsx(self):
        self.assertEqual(self.a.usuario.plan, 'COMMON')
        self.assertEqual(self.client.get(EXPORTAR_TRANSACOES).status_code, 403)

        resposta = self.client.get(EXPORTAR)

        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(
            resposta['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertIn('attachment;', resposta['Content-Disposition'])
        self.assertEqual(self.planilha(resposta).sheetnames, ABAS)

    def test_cada_aba_traz_so_os_dados_do_proprio_usuario(self):
        planilha = self.planilha(self.client.get(EXPORTAR))

        esperados = {
            'Transações': ('Feira do mês Usuário A', 'Feira do mês Usuário B'),
            'Contas': ('Conta A', 'Conta B'),
            'Cartões': ('Cartão A', 'Cartão B'),
            'Metas': ('Meta A', 'Meta B'),
            'Orçamentos': ('Mercado A', 'Mercado B'),
        }
        for aba, (do_dono, do_outro) in esperados.items():
            with self.subTest(aba=aba):
                textos = self.textos(planilha[aba])
                self.assertIn(do_dono, textos)
                self.assertNotIn(do_outro, textos)
        self.assertIn('- 87.50', self.textos(planilha['Transações']))
        self.assertIn('300.00', self.textos(planilha['Orçamentos']))

    def test_sem_login_recebe_401(self):
        self.client.force_authenticate(user=None)

        self.assertEqual(self.client.get(EXPORTAR).status_code, 401)
